import sqlite3 as lite
from io import StringIO, BytesIO
from fpdf import FPDF

from docx import Document
import pdfkit
from htmldocx import HtmlToDocx
from lighthouse import *


class Stat:
    def __init__(self, db_path):
        self.con = lite.connect(db_path)
        self.cur = self.con.cursor()

    def __query(self, num1, num2, table_name):
        return '''
        select x.number, out_y.out_n, in_y.in_n
        from (
        select distinct x.number
        from (
        select {0} as number
        from {2}
        union
        select {1} as number
        from {2}
        ) x 
        ) x

        join
        (select {0}, count({0}) as out_n
        from {2}
        group by {0}
        ) out_y
        on x.number=out_y.num1

        join
        (select {1}, count({1}) as in_n
        from {2}
        group by {1}
        ) in_y
        on x.number=in_y.num2
        '''.format(num1, num2, table_name)

    def __get_data(self, num1, num2, table):
        data = []
        for row in self.cur.execute(self.__query(num1, num2, table)):
            data.append(row)
        return data

    def __tables(self):
        table_items = self.cur.execute("SELECT name FROM sqlite_master WHERE type='table' and name='table_items'").fetchone()

        tables = []
        for id in self.cur.execute("SELECT id FROM {0} WHERE header_id='3cb7bbc7-2bc4-4987-9b6b-1f533b4f174e'".format(table_items[0])):
            for table in self.cur.execute("SELECT name FROM sqlite_master WHERE type='table' and name like '%{0}'".format(id[0].replace('-', ''))):
                tables.append(table)

        return tables

    def reports_html(self, num1, num2):
        files = []

        header = '''
        <!DOCTYPE html>
        <html>
        <head>
        <style>
            table, th, td { border: 1px solid black; }
        </style>
        </head>
        <body>
                '''
        table_header = '''
            <br>
            <table style="width:100%">
                <tr>
                    <th>Номер</th>
                    <th>Количество исходящих</th> 
                    <th>Количество входящих</th>
                </tr>  
                '''
        table_footer = '\n\t</table>'
        footer = '''
        </body>
        </html>
                '''
        body = ''

        for table in self.__tables():
            body = body + table_header
            data = self.__get_data(num1, num2, table[0])
            for row in data:
                body = body + '<tr>'
                body = body + '\n' + '\t\t<td>' + row[0] + '</td>'
                body = body + '\n' + '\t\t<td>' + str(row[1]) + '</td>'
                body = body + '\n' + '\t\t<td>' + str(row[2]) + '</td>'
                body = body + '\n' + '\t</tr>\n\t'
            body = body + table_footer

            files.append(['{0}.html'.format(table[0]), header + body + footer])

        return files

    def reports_docx(self, num1, num2):
        files = []

        for table in self.__tables():
            document = Document()

            data = self.__get_data(num1, num2, table[0])
            doc_table = document.add_table(rows=len(data) + 1, cols=3)
            # table.style = 'LightShading-Accent1'
            row = 0
            doc_table.cell(row, 0).text = 'Номер'
            doc_table.cell(row, 1).text = 'Исходящие'
            doc_table.cell(row, 2).text = 'Входящие'

            row = row + 1
            for data_row in data:
                doc_table.cell(row, 0).text = data_row[0]
                doc_table.cell(row, 1).text = str(data_row[1])
                doc_table.cell(row, 2).text = str(data_row[2])
                row = row + 1

            stream = BytesIO()
            document.save(stream)
            files.append(['{0}.docx'.format(table[0]), bytearray(stream.getbuffer())])

        return files

    def reports_pdf(self, num1, num2):
        pdf = FPDF(format='a4', unit='cm')
        pdf.add_page()
        pdf.set_font('Times', '', 10.0)

        # pdf.add_font('DejaVu', '', 'DejaVuSansCondensed.ttf', uni=True)
        # pdf.set_font('DejaVu', '', 14)

        epw = pdf.w - 2 * pdf.l_margin
        col_width = epw / 4
        th = pdf.font_size

        files = []
        for table in self.__tables():
            data = self.__get_data(num1, num2, table[0])
            pdf.cell(col_width, th, str('#'), border=1)
            pdf.cell(col_width, th, str('->'), border=1)
            pdf.cell(col_width, th, str('<-'), border=1)
            pdf.ln(th)

            for data_row in data:
                pdf.cell(col_width, th, str(data_row[0]), border=1)
                pdf.cell(col_width, th, str(data_row[1]), border=1)
                pdf.cell(col_width, th, str(data_row[2]), border=1)
                pdf.ln(th)

            files.append(['{0}.pdf'.format(table[0]), pdf.output(dest='S').encode('latin1')])

        return files

    def reports(self, num1, num2):
        resources = []

        files_html = self.reports_html(num1, num2)
        files_docx = self.reports_docx(num1, num2)
        files_pdf = self.reports_pdf(num1, num2)

        index = 0
        for _ in self.__tables():
            files = []
            files.append([files_html[index][0], files_html[index][1]])
            files.append([files_docx[index][0], files_docx[index][1]])
            files.append([files_pdf[index][0], files_pdf[index][1]])

            resources.append(files)
            index = index + 1

        return resources

class StatHeader(metaclass=Header):
    display_name = 'Stat'
    resource_html = Field('HTML', ValueType.String, binary_type=BinaryType.Resource)
    resource_docx = Field('DOCX', ValueType.String, binary_type=BinaryType.Resource)
    resource_pdf = Field('PDF', ValueType.String, binary_type=BinaryType.Resource)

class SqliteTask(Task):

    def get_id(self):
        return 'b795762d-6a93-4bc6-8a42-6440d400107f'

    def get_category(self):
        return 'StatTask'

    def get_display_name(self):
        return 'Stat Task'

    def get_headers(self):
        return HeaderCollection(StatHeader)

    def get_enter_params(self):
        return EnterParamCollection(
            EnterParamField('path', 'Path', ValueType.String, is_array=False, required=True, file_path=True,
                            default_value='', category='Required', description='Path'))

    @staticmethod
    def __fill_result(result_writer, resources):
        for resource in resources:
            line = StatHeader.create_empty()
            line[StatHeader.resource_html] = Resource(resource[0][0], resource[0][1].encode())
            line[StatHeader.resource_docx] = Resource(resource[1][0], resource[1][1])
            line[StatHeader.resource_pdf] = Resource(resource[2][0], resource[2][1])
            # line[StatHeader.resource_html] = Resource(resource[0][0], resource[0][1])
            # line[StatHeader.resource_docx] = Resource(resource[1][0], resource[1][1])
            # line[StatHeader.resource_pdf] = Resource(resource[2][0], resource[2][1])

            result_writer.write_line(line, header_class=StatHeader)

    def execute(self, enter_params, result_writer, log_writer, temp_directory):
        stat = Stat(enter_params.path)
        self.__fill_result(result_writer, stat.reports('num1', 'num2'))

# for tests
# stat = Stat('D:\\dev\\nt\\lampyre\\master\\AS\\AS.UV\\bin\\Debug\\net472\\!sqlite_project\\3b5cbesa.rk3\\ca6b84c20b7644a3b9fbf6170876f3bf.sqlite')
# resources = stat.reports('num1', 'num2')
#
# for resource in resources:
#     print(resource[0][0])
#     print(resource[1][0])
#     print(resource[2][0])
#
# file = open("d:/temp/" + resources[0][0][0], "w")
# file.write(resources[0][0][1])
# file.close()
#
# file = open("d:/temp/" + resources[0][1][0], "wb")
# file.write(resources[0][1][1])
# file.close()
#
# file = open("d:/temp/" + resources[0][2][0], "wb")
# file.write(resources[0][2][1])
# file.close()
